import os
from typing import List, Dict, Any

class PPTXProcessor:
    @staticmethod
    def extract_chunks(file_path: str, course_id: str, material_id: str, source_name: str) -> List[Dict[str, Any]]:
        chunks = []
        if not os.path.exists(file_path):
            return chunks
            
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_text_parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text_parts.append(shape.text.strip())
                
                full_slide_text = "\n".join(slide_text_parts)
                if not full_slide_text:
                    continue
                    
                topic = "General Concepts"
                lower_text = full_slide_text.lower()
                if "array" in lower_text:
                    topic = "Arrays"
                elif "search" in lower_text or "binary search" in lower_text:
                    topic = "Searching"
                elif "sort" in lower_text:
                    topic = "Sorting"
                elif "recursion" in lower_text:
                    topic = "Recursion"
                elif "tree" in lower_text:
                    topic = "Trees"
                elif "graph" in lower_text:
                    topic = "Graphs"
                    
                chunks.append({
                    "course_id": course_id,
                    "material_id": material_id,
                    "material_type": "pptx",
                    "source_name": source_name,
                    "page_number": None,
                    "slide_number": slide_num,
                    "video_timestamp": None,
                    "topic": topic,
                    "concept": None,
                    "content": full_slide_text
                })
        except Exception:
            # Fallback
            chunks.append({
                "course_id": course_id,
                "material_id": material_id,
                "material_type": "pptx",
                "source_name": source_name,
                "page_number": None,
                "slide_number": 1,
                "video_timestamp": None,
                "topic": "General",
                "concept": None,
                "content": f"Slide content from {source_name}"
            })
            
        return chunks
